# -*- coding: utf-8 -*-
"""视频类型 —— 把表格 I 列那一个词，翻译成描述语模板 + 运镜卡池。

═══ 这一层解决什么 ═══

A 层描述语结尾的机读标签【闭眼|常规】是 Gemini 看图猜的。猜「闭没闭眼」很准，
猜「这张图到底有没有主角」不准 —— 0826 那批 v19~v27 是一群人低头一起祷告，
AI 每一张都硬挑一个人当主角、让他对口型，镜头再推成他的面部特写，9 个编号全部穿帮。

这不是 AI 蠢。同一张十几个人围成一圈的图，要拍成「其中一个人在祷告」还是
「一群人在祷告」，是创作决定，不是识别问题 —— 该由制作人定。

所以流程变成：

    制作人看图 → 表格 I 列写类型 → 扫表落成 当天文件夹/类型登记.json
                                        ↓
              描述语生成器.html 按类型切写法（主角版 / 群像版）
                                        ↓
              flow_shots 按类型给的场景标签挑运镜卡池
                                        ↓
              提交脚本拼提示词

**I 列留空 = 普通祷告 = 老行为一个字都不变。** 这是这个模块最重要的性质：
不填就等于这套东西不存在，绝不会把已经跑顺的批次搞乱。

═══ 类型登记.json ═══

落在当天图片文件夹里，格式极简，人也能直接改：

    {"v19": "群体祷告", "v20": "群体祷告", "v31": "车内祷告"}

之所以要落成文件而不是每次去查表格：下游四五个脚本都要用，
每个都去连一次 Google 表格既慢又多一个失败点；而且接力/补拍文件夹是复制出去的，
带着这个文件走，到哪都还认得。
"""

from __future__ import annotations

import json
import os
import re

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
TYPES_FILE = os.path.join(SCRIPT_DIR, "video_types.json")
REGISTRY_NAME = "类型登记.json"

_cache = None

# A 层描述语结尾的机读标签，和 flow_shots.TAG_RE 保持同一套词表
_TAG_RE = re.compile(r"【\s*(闭眼|睁眼|空镜)\s*[|｜]\s*(常规|舞台|车内|围绕|群像)\s*】\s*$")


def load_types(force=False):
    global _cache
    if _cache is None or force:
        try:
            with open(TYPES_FILE, "r", encoding="utf-8") as f:
                _cache = json.load(f)
        except Exception:
            _cache = {"默认类型": "普通祷告", "类型": {}}
    return _cache


def default_type():
    return load_types().get("默认类型") or "普通祷告"


def all_types():
    return list((load_types().get("类型") or {}).keys())


def normalize(raw):
    """I 列里那个词 → 正式类型名。认不出来返回 None（调用方自己决定要不要报警）。

    留空 / 全空白 → 默认类型（普通祷告），这是最常见的情况。
    """
    s = re.sub(r"\s+", "", str(raw or ""))
    s = s.replace("（", "(").replace("）", ")")
    types = load_types().get("类型") or {}
    if not s:
        return default_type()
    for name, spec in types.items():
        if s == re.sub(r"\s+", "", name):
            return name
        for a in (spec.get("别名") or []):
            if a and s == re.sub(r"\s+", "", str(a)):
                return name
    # 再宽松一层：写「群体」「群像祷告」「这是群体祷告」都认
    for name, spec in types.items():
        keys = [name] + [a for a in (spec.get("别名") or []) if a]
        for k in keys:
            k = re.sub(r"\s+", "", str(k))
            if len(k) >= 2 and (k in s or s in k):
                return name
    return None


def spec_of(type_name):
    return (load_types().get("类型") or {}).get(type_name) or {}


def tag_of(type_name):
    """这个类型强制写死的机读标签，如 "空镜|群像"。None = 不强制，交给 AI 判。"""
    return spec_of(type_name).get("标签")


def template_of(type_name):
    return spec_of(type_name).get("描述语模板") or "主角版"


def eye_scene_of(type_name, fallback=("闭眼", "常规")):
    """类型 → (眼睛, 场景)。类型不强制标签时返回 fallback（一般是 AI 判出来的那个）。"""
    tag = tag_of(type_name)
    if not tag:
        return fallback
    parts = re.split(r"[|｜]", str(tag))
    if len(parts) != 2:
        return fallback
    return parts[0].strip(), parts[1].strip()


# ═══════════════ 类型登记.json 读写 ═══════════════

def registry_path(folder):
    return os.path.join(folder or ".", REGISTRY_NAME)


def norm_label(x):
    m = re.match(r"^\s*[vV]?0*(\d+)\s*$", str(x or ""))
    return "v%02d" % int(m.group(1)) if m else str(x or "").strip()


def load_registry(folder):
    """{v19: "群体祷告"}。没有这个文件就返回空 dict —— 那就全走默认类型。"""
    try:
        with open(registry_path(folder), "r", encoding="utf-8-sig") as f:
            raw = json.load(f)
    except Exception:
        return {}
    if not isinstance(raw, dict):
        return {}
    out = {}
    for k, v in raw.items():
        if str(k).startswith("_"):
            continue
        name = normalize(v)
        if name:
            out[norm_label(k)] = name
    return out


def save_registry(folder, mapping, note=""):
    """写登记表。带一段人话说明，方便手工改。"""
    body = {"_说明": ("制作人在表格 I 列填的视频类型，扫表时落到这里。"
                      "手工改这个文件也生效。可用类型：" + "、".join(all_types())
                      + "。删掉某个编号 = 该编号回到默认类型（"
                      + default_type() + "）。" + (note or ""))}
    for k in sorted(mapping, key=lambda x: (len(x), x)):
        body[norm_label(k)] = mapping[k]
    tmp = registry_path(folder) + ".tmp"
    with open(tmp, "w", encoding="utf-8") as f:
        json.dump(body, f, ensure_ascii=False, indent=1)
        f.flush()
        os.fsync(f.fileno())
    os.replace(tmp, registry_path(folder))
    return registry_path(folder)


def type_of(folder, label, registry=None):
    reg = load_registry(folder) if registry is None else registry
    return reg.get(norm_label(label)) or default_type()


# ═══════════════ 把类型盖到描述语的机读标签上 ═══════════════

def force_tag(desc, type_name):
    """按类型改写描述语结尾的机读标签。制作人说了算，AI 判的只是兜底。

    类型不强制标签（普通祷告）时原样返回。
    描述语本来没有标签的，直接补一个上去。
    """
    tag = tag_of(type_name)
    if not tag or not (desc or "").strip():
        return desc
    want = "【%s】" % tag.replace("｜", "|")
    lines = desc.rstrip().splitlines()
    if not lines:
        return desc
    first = lines[0].rstrip()
    if _TAG_RE.search(first):
        first = _TAG_RE.sub("", first).rstrip()
    lines[0] = first + want
    return "\n".join(lines)


def voice_of(desc):
    """描述语里的【音色】行 → 音色锁用的那串。没有返回空串。

    群像（画外音）尤其需要它：画面里没人说话，解说是谁的声音只能靠这里锁。
    """
    for line in (desc or "").splitlines():
        s = line.strip()
        if s.startswith("【音色】"):
            return s[len("【音色】"):].strip()
    return ""


# ═══════════════ 从表格 I 列扫一遍 ═══════════════

def scan_sheet(folder, cfg=None, type_col="I", log=print):
    """连 Google 表格读 I 列，写 类型登记.json。

    编号怎么来、哪些行算数，全部照抄 表格自动下载器.fetch_sheet_rows 的规矩：
    优先 name_col 里的 v 号，没有就按候选行序号补；B 列过滤放在编号之后
    （否则跳过 Tiktok 行会让 v 号整体平移，对不上图）。鉴权和分表选择也直接
    复用那边的函数，不另起一套 —— 两处规则一旦漂移，扫出来的类型就会挂错编号。

    拿不到表格（没网 / 没授权 / 表格结构变了）不报错退出，返回 None：
    下游继续走默认类型（普通祷告 = 老行为），这条链路不该成为出片的阻塞点。
    """
    cfg = cfg or {}
    try:
        import openpyxl
        import sheet_filter
        import 表格自动下载器 as TB
    except Exception as e:
        log(f"  ! 读不到 表格自动下载器 / openpyxl（{str(e)[:60]}），跳过扫 I 列")
        return None

    url = cfg.get("sheet_url")
    if not url:
        log("  ! 配置里没有 sheet_url，跳过扫 I 列")
        return None
    name_col = (cfg.get("name_column") or "N").strip().upper()
    link_col = (cfg.get("link_column") or "J").strip().upper()
    kind_col = (cfg.get("type_column") or "B").strip().upper()
    type_col = (type_col or "I").strip().upper()
    mode = cfg.get("row_filter") or "normal"
    blank_as = cfg.get("blank_as") or "normal"

    tmp = os.path.join(SCRIPT_DIR, "_types_tmp.xlsx")
    try:
        token = None
        try:
            token = TB.try_get_drive_token(cfg, log)
        except Exception:
            pass
        content = TB.download_sheet_xlsx(url, log, token=token)
        with open(tmp, "wb") as f:
            f.write(content)
        wb = openpyxl.load_workbook(tmp)
        ws = TB.pick_worksheet(wb, url, cfg.get("sheet_name", ""), log)
    except Exception as e:
        log(f"  ! 拉表格失败（{str(e)[:80]}），跳过扫 I 列")
        return None
    finally:
        pass

    mapping, unknown = {}, []
    seq = 0
    try:
        for i in range(1, ws.max_row + 1):
            cell = ws[f"{link_col}{i}"]
            u = cell.hyperlink.target if cell.hyperlink else None
            text = str(cell.value).strip() if cell.value else ""
            if not u and text.startswith("http"):
                u = text
            label = None
            nv = ws[f"{name_col}{i}"].value
            if nv:
                m = re.search(r"[vV](\d+)", str(nv))
                if m:
                    label = "v%02d" % int(m.group(1))
            if not u and not label:
                continue
            seq += 1
            if not label:
                label = "v%02d" % seq
            kind, _extra = sheet_filter.row_kind(ws[f"{kind_col}{i}"].value)
            if not sheet_filter.wants(kind, mode, blank_as=blank_as):
                continue
            raw = ws[f"{type_col}{i}"].value
            name = normalize(raw)
            if name is None:
                unknown.append((label, str(raw)[:24]))
                continue
            mapping[label] = name
    finally:
        try:
            os.remove(tmp)
        except OSError:
            pass

    if unknown:
        log("  ! I 列这些词不认识，按默认类型处理："
            + "、".join(f"{a}={b}" for a, b in unknown[:8]))
    save_registry(folder, mapping)
    from collections import Counter
    c = Counter(mapping.values())
    log("  ✓ 类型登记 %d 条 → %s：%s"
        % (len(mapping), registry_path(folder),
           "，".join(f"{k}×{v}" for k, v in c.items())))
    return mapping


# ═══════════════ 类型 × 描述语写法 体检 ═══════════════

# 主角版描述语的指纹：只要出现其中一条，这段描述语就是围着某一个人写的。
_主角痕迹 = ("这是一位", "他正在", "她正在", "男性基督徒", "女性基督徒",
             "全程紧闭双眼", "全程睁着眼", "双脚不移动位置", "有声地祷告",
             "他以", "她以", "保持同一人物")


def desc_conflicts(desc, type_name):
    """描述语的写法和类型对不对得上。对得上返回空串，对不上返回一句人话。

    只查一件事，但这件事最要命：类型是群体祷告（群像版），
    描述语却还是主角版 —— 提示词会自相矛盾，模型挑错边就穿帮。
    反过来（类型是普通祷告、描述语是群像版）不查：那是制作人主动降级，
    没有矛盾，只是镜头保守一点。
    """
    if template_of(type_name) != "群像版":
        return ""
    body = str(desc or "")
    hits = [w for w in _主角痕迹 if w in body]
    if not hits:
        return ""
    return ("类型是「%s」（群像版），但描述语还是主角版写法：出现了 %s。"
            "请用 描述语生成器.html 对这张图重跑一次描述语。"
            % (type_name, "、".join(hits[:3])))


# ═══════════════ 场景标签体检：围绕 × 只有一两个人 ═══════════════

# 主角之外真有一群人时，描述语里必然出现这些词之一。
_人群词 = re.compile(r"众人|一圈|围成|人群|十几位|二十多位|多位|几位|一群|"
                     r"满场|手拉手|按手|簇拥|合围|信徒们|会众")
# 「众人」单独拎出来：0829 那批每一张都只有这一个词，而且是被提示词逼出来的
# 那句「被众人围在中间」——它不能算作"画面里真有一群人"的证据。
_空头众人 = re.compile(r"被众人围在中间")


def scene_conflicts(desc):
    """描述语【自己内部】的矛盾（和表格类型无关）。对得上返回空串。

    只查一件事，但它 0829 让整批 10 条白跑：结尾标签写了【围绕】，
    正文却是一个坐在演播室扶手椅里的人、画面里只有一两个人。

    为什么会这样：v5 及以前的提示词在「睁眼」这一档只给了两个选择 ——
    大幅度来回走动，或者「被众人围在中间、双脚不离开原地」。坐着的人显然
    走不了，模型只能选第二句，于是凭空写出一圈根本不存在的人，标签跟着
    判成围绕。B 层再据此附加「人群保留」和「贴着人群外缘绕行」的运镜，
    画面里没有那圈人，模型只能造人或者拍空 —— 每一条都废。
    （A 层已在 描述语生成器 v6 补上「坐着」这一档，这里是事后那道网。）

    两条判据，命中任一条就报：
      · 围绕 + 主角是坐着的（围圈祷告的人都站着；坐着的是访谈，不是围绕）
      · 围绕 + 正文里除了那句「被众人围在中间」之外找不到任何一群人的描写
    """
    body = str(desc or "").strip()
    if not body:
        return ""
    first = body.splitlines()[0].rstrip()
    m = _TAG_RE.search(first)
    if not m or m.group(2) != "围绕":
        return ""
    try:
        from flow_shots import is_seated
        seated = is_seated(first)
    except Exception:
        seated = bool(re.search(r"以[^，,。；;]{0,8}坐姿|端坐|(?:她|他|正)坐在",
                                first))
    if seated:
        return ("标签写的是【围绕】，但正文里人物是坐着的 —— 围圈祷告的人都站着，"
                "坐着说话的多半是访谈图，应该是【常规】。照围绕跑会附加"
                "「人群保留 + 贴着人群外缘绕行」，画面里没有那圈人，模型只能凭空造人。")
    rest = _空头众人.sub("", first)
    if not _人群词.search(rest):
        return ("标签写的是【围绕】，但正文里除了「被众人围在中间」这句套话之外，"
                "找不到任何一群人的描写 —— 多半是画面里只有一两个人被硬写成了围绕。")
    return ""


def audit(folder, labels=None, log=print):
    """整批体检：哪些编号的描述语和类型对不上。返回 [(编号, 原因)]。"""
    from desc_files import load_desc_map
    reg = load_registry(folder)
    descs = load_desc_map(folder, log=lambda *a: None)
    bad = []
    for lab in sorted(labels or descs, key=lambda x: (len(x), x)):
        t = reg.get(norm_label(lab)) or default_type()
        d = descs.get(norm_label(lab)) or ""
        # 两道体检：① 类型 × 写法（群像版 vs 主角版）② 描述语自己内部的矛盾
        why = desc_conflicts(d, t) or scene_conflicts(d)
        if why:
            bad.append((norm_label(lab), why))
    if bad and log:
        log("⚠ %d 个编号的描述语有问题（类型对不上，或标签和正文自相矛盾）：" % len(bad))
        for lab, why in bad[:12]:
            log("    %s  %s" % (lab, why))
        if len(bad) > 12:
            log("    …… 其余 %d 个同理" % (len(bad) - 12))
    return bad


def load_voice_map(folder):
    """{v01: "白人|男|中年|中低音|美国口音"}，直接读 vNN_desc.txt 原文。

    为什么不用 desc_files.load_desc_map：那个函数返回的是 desc_body()，
    已经把【音色】行剥掉了（提交 A 层身份锁时确实不该带这一行）。
    音色要单独读原文，否则 voice_lock 永远是空的 —— 音色锁那一块也就永远不生效。
    """
    import re as _re
    out = {}
    if not folder or not os.path.isdir(folder):
        return out
    for name in os.listdir(folder):
        m = _re.match(r"^(v\d+)_desc\.txt$", name, _re.I)
        if not m:
            continue
        try:
            with open(os.path.join(folder, name), "r", encoding="utf-8-sig") as f:
                raw = f.read()
        except OSError:
            continue
        v = voice_of(raw)
        if v:
            out[norm_label(m.group(1))] = v
    return out
